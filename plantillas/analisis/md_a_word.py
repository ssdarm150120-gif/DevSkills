"""Convierte los documentos Markdown del proyecto a Word (.docx).

Uso:  python documentacion/analisis/md_a_word.py [--mermaid CARPETA] [--solo FILTRO] [--proyecto NOMBRE]

Genera documentacion/word/*.docx a partir de SRS.md, DISENO.md, PLAN.md, ui/prompts_figma_make.md,
EVALUACION_ALTERNATIVAS.md y cualquier otro .md de primer nivel de documentacion/ (salvo README.md, CONTEXTO.md y las
plantillas sin llenar, que tienen `<!--` en el título).
Los .md son la fuente de verdad: si cambia un requisito, se edita el .md y se regenera el Word.

- `--mermaid CARPETA`: carpeta donde instalaste mermaid-cli (`npm install --save-exact @mermaid-js/mermaid-cli@12.0.0`),
  con `node_modules/.bin/mmdc` y, si hace falta, un `puppeteer.json` (por ejemplo, para usar el Chrome
  local). Sin esta opción se usa `mmdc` si está en el PATH. Sin mermaid-cli, los diagramas quedan
  como bloque de código y el script lo avisa. Los PNG se guardan en caché en `word/_diagramas/`.
- `--solo FILTRO`: convierte solo los documentos cuya ruta contiene FILTRO (por ejemplo, `SRS`).
- `--proyecto NOMBRE`: nombre del proyecto para la portada, el encabezado y los archivos. Por defecto
  se lee del segundo encabezado del SRS (`## <Nombre del proyecto>`).

Requiere python-docx y Pillow (`pip install python-docx==1.2.0 pillow==12.3.0`), con el Python del sistema o un
entorno propio de la documentación, no el entorno del código.
"""
import argparse, hashlib, re, shutil, subprocess, pathlib, sys, unicodedata
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# En Windows la consola y las tuberías usan la página de códigos local (cp1252): sin esto los acentos
# salen corruptos. La salida de DevSkills siempre es UTF-8.
for _flujo in (sys.stdout, sys.stderr):
    if hasattr(_flujo, "reconfigure"):
        _flujo.reconfigure(encoding="utf-8")

DOC = pathlib.Path(__file__).resolve().parent.parent  # carpeta documentacion/
SALIDA = DOC / "word"
CACHE = SALIDA / "_diagramas"
MMDC = None            # comando de mermaid-cli (se resuelve en main)
PUPPETEER = None       # puppeteer.json opcional
PROYECTO = "Proyecto"  # se resuelve en main
AVISO_MERMAID = []     # para avisar una sola vez

PETROLEO = RGBColor(0x0F, 0x4C, 0x5C)
GRIS = RGBColor(0x47, 0x55, 0x69)
ANCHO_UTIL = Cm(17.0)

# (fuente, prefijo del archivo Word, título para el encabezado). El nombre final es <prefijo>_<Proyecto>.docx
DOCUMENTOS = [
    ("SRS.md", "SRS", "Especificación de Requisitos de Software"),
    ("DISENO.md", "Diseno_Tecnico", "Documento de Diseño de Software"),
    ("PLAN.md", "Plan_de_Entregas", "Plan de entregas ágil"),
    ("ui/prompts_figma_make.md", "Prompts_Figma_Make", "Prompts para Figma Make"),
    ("EVALUACION_ALTERNATIVAS.md", "Evaluacion_de_Alternativas", "Evaluación de alternativas"),
]


def nombre_archivo(texto):
    """'Mi Proyecto Ágil' -> 'Mi_Proyecto_Agil' (sin acentos ni símbolos)."""
    t = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Za-z0-9]+", "_", t).strip("_") or "Proyecto"


def sin_llenar(md):
    """Una plantilla sin llenar (con `<!--` en el título) no se convierte."""
    h1 = next((l for l in md.read_text(encoding="utf-8").split("\n") if l.startswith("# ")), "")
    return "<!--" in h1


def documentos():
    """Los documentos fijos que existan, más cualquier otro .md de primer nivel de documentacion/."""
    lista = [d for d in DOCUMENTOS if (DOC / d[0]).exists() and not sin_llenar(DOC / d[0])]
    fijos = {d[0] for d in DOCUMENTOS}
    for md in sorted(DOC.glob("*.md")):
        if md.name not in fijos and md.name.upper() not in ("README.MD", "CONTEXTO.MD") and not sin_llenar(md):
            h1 = next((l[2:].strip() for l in md.read_text(encoding="utf-8").split("\n") if l.startswith("# ")), md.stem)
            lista.append((md.name, nombre_archivo(md.stem.title()), h1))
    return lista

# ---------------------------------------------------------------- utilidades XML
def sombrear(celda, color):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
    tcPr.append(shd)

def borde_parrafo(p, lado="bottom", color="0F4C5C", tam="8"):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement(f"w:{lado}")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), tam); b.set(qn("w:space"), "4"); b.set(qn("w:color"), color)
    bdr.append(b); pPr.append(bdr)

def sombrear_parrafo(p, color):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
    pPr.append(shd)

def campo(parrafo, instruccion):
    r = parrafo.add_run()
    for tipo, texto in (("begin", None), (None, instruccion), ("separate", None), (None, "Actualiza la tabla: clic derecho → Actualizar campo"), ("end", None)):
        if tipo:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), tipo); r._r.append(fc)
        elif texto == instruccion:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instruccion; r._r.append(it)
        else:
            t = OxmlElement("w:t"); t.text = texto; r._r.append(t)

def actualizar_campos_al_abrir(doc):
    s = doc.settings.element
    u = OxmlElement("w:updateFields"); u.set(qn("w:val"), "true"); s.append(u)

# ---------------------------------------------------------------- texto en línea
TOKEN = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|(?<![\w*])\*[^*\s][^*]*?\*(?![\w*]))")

def agregar_texto(p, texto, tam=None, color=None, negrita=False):
    texto = texto.replace("<br>", "\n").replace("\\|", "|")
    for i, linea in enumerate(texto.split("\n")):
        if i:
            p.add_run().add_break()
        for parte in TOKEN.split(linea):
            if not parte:
                continue
            if parte.startswith("**") and parte.endswith("**"):
                r = p.add_run(parte[2:-2]); r.bold = True
            elif parte.startswith("`") and parte.endswith("`"):
                r = p.add_run(parte[1:-1]); r.font.name = "Consolas"; r.font.color.rgb = RGBColor(0x9F, 0x12, 0x39)
            elif parte.startswith("[") and "](" in parte:
                r = p.add_run(parte[1:parte.index("](")]); r.underline = True; r.font.color.rgb = PETROLEO
            elif parte.startswith("*") and parte.endswith("*") and len(parte) > 2:
                r = p.add_run(parte[1:-1]); r.italic = True
            else:
                r = p.add_run(parte)
            if negrita:
                r.bold = True
            if tam:
                r.font.size = tam
            if color is not None and not (parte.startswith("`")):
                r.font.color.rgb = color

# ---------------------------------------------------------------- bloques
def tabla(doc, filas):
    filas = [[c.strip() for c in f.strip().strip("|").split("|")] for f in filas]
    enc, cuerpo = filas[0], [f for f in filas[2:]]
    n = len(enc)
    cuerpo = [(f + [""] * n)[:n] for f in cuerpo]
    # anchos proporcionales al contenido (con mínimo)
    # ancho mínimo: la palabra más larga de la columna (para no partir palabras letra por letra)
    def palabra_max(txt):
        return max((len(w) for w in re.sub(r"[*`]", "", txt).split()), default=1)
    minimo = [min(18, max(4, max([palabra_max(enc[i])] + [palabra_max(f[i]) for f in cuerpo]) + 1)) for i in range(n)]
    largo = [max([len(enc[i])] + [min(len(f[i]), 120) for f in cuerpo]) for i in range(n)]
    tot_min = sum(minimo)
    extra = [max(0, l - m) for l, m in zip(largo, minimo)]
    tot_extra = sum(extra) or 1
    caracteres = 105  # caracteres aproximados que entran en 17 cm a 8 pt
    libre = max(0, caracteres - tot_min)
    unidades = [m + libre * e / tot_extra for m, e in zip(minimo, extra)]
    tot = sum(unidades)
    anchos = [int(ANCHO_UTIL * u / tot) for u in unidades]
    t = doc.add_table(rows=1 + len(cuerpo), cols=n)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tam = Pt(8) if n >= 4 else Pt(9)
    for j, txt in enumerate(enc):
        c = t.rows[0].cells[j]; c.width = anchos[j]; sombrear(c, "0F4C5C")
        p = c.paragraphs[0]; agregar_texto(p, txt, tam, RGBColor(0xFF, 0xFF, 0xFF), negrita=True)
    trPr = t.rows[0]._tr.get_or_add_trPr(); h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trPr.append(h)
    for i, f in enumerate(cuerpo, 1):
        for j, txt in enumerate(f):
            c = t.rows[i].cells[j]; c.width = anchos[j]
            if i % 2 == 0:
                sombrear(c, "F1F5F9")
            agregar_texto(c.paragraphs[0], txt, tam)
    doc.add_paragraph()

def codigo(doc, lineas, lenguaje):
    t = doc.add_table(rows=1, cols=1); t.style = "Table Grid"; t.autofit = False
    c = t.rows[0].cells[0]; c.width = ANCHO_UTIL; sombrear(c, "F8FAFC")
    p = c.paragraphs[0]
    claves = {"Característica:", "Escenario:", "Dado", "Cuando", "Entonces", "Y", "Pero", "Mientras", "Antes", "Porque"}
    for k, l in enumerate(lineas):
        if k:
            p.add_run().add_break()
        prim = l.strip().split(" ")[0] if l.strip() else ""
        r = p.add_run(l)
        r.font.name = "Consolas"; r.font.size = Pt(7.5)
        if lenguaje == "gherkin" and prim in claves:
            r.bold = True; r.font.color.rgb = PETROLEO
    doc.add_paragraph()

def mermaid(doc, lineas):
    if MMDC is None:
        if not AVISO_MERMAID:
            AVISO_MERMAID.append(1)
            print("Aviso: no encontré mermaid-cli (usa --mermaid CARPETA o pon mmdc en el PATH); "
                  "los diagramas Mermaid quedan como bloque de código.")
        return codigo(doc, lineas, "mermaid")
    fuente = "\n".join(lineas)
    h = hashlib.sha256(fuente.encode("utf-8")).hexdigest()[:12]  # clave de caché, no de seguridad
    png = CACHE / f"{h}.png"
    if not png.exists():
        CACHE.mkdir(parents=True, exist_ok=True)
        src = CACHE / f"{h}.mmd"; src.write_text(fuente, encoding="utf-8")
        cmd = [MMDC] + (["-p", str(PUPPETEER)] if PUPPETEER else []) + ["-i", str(src), "-o", str(png), "-s", "3", "-b", "white"]
        # nosemgrep: python.lang.security.audit.dangerous-subprocess-use-tainted-env-args.dangerous-subprocess-use-tainted-env-args -- lista de argumentos, sin shell, con ejecutable y rutas constantes
        res = subprocess.run(cmd, capture_output=True, timeout=180)
        src.unlink(missing_ok=True)
        if res.returncode != 0:
            print(f"Aviso: mermaid-cli no pudo renderizar un diagrama ({h}); queda como bloque de código.")
    if not png.exists():
        return codigo(doc, lineas, "mermaid")
    from PIL import Image
    w, hgt = Image.open(png).size
    ancho = min(ANCHO_UTIL, Cm(16.5))
    alto = ancho * hgt / w
    if alto > Cm(21):
        ancho = int(ancho * Cm(21) / alto)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(png), width=int(ancho))

# ---------------------------------------------------------------- documento
def preparar(doc, titulo_doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for m in ("left_margin", "right_margin"):
        setattr(sec, m, Cm(2))
    sec.top_margin = sec.bottom_margin = Cm(2)
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.paragraph_format.space_after = Pt(4)
    for n, tam in ((1, 16), (2, 13), (3, 11.5), (4, 10.5)):
        h = doc.styles[f"Heading {n}"]; h.font.name = "Calibri"; h.font.size = Pt(tam)
        h.font.color.rgb = PETROLEO; h.font.bold = True
        h.paragraph_format.space_before = Pt(12 if n <= 2 else 8); h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
    # encabezado y pie
    enc = sec.header.paragraphs[0]; enc.text = f"{PROYECTO} · {titulo_doc}"
    enc.runs[0].font.size = Pt(8); enc.runs[0].font.color.rgb = GRIS
    pie = sec.footer.paragraphs[0]; pie.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = pie.add_run("Página "); r.font.size = Pt(8)
    campo_pag = pie.add_run()
    for tipo, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if tipo:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), tipo); campo_pag._r.append(fc)
        else:
            it = OxmlElement("w:instrText"); it.text = txt; campo_pag._r.append(it)
    campo_pag.font.size = Pt(8)

def portada(doc, titulo, subtitulo):
    for _ in range(8):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(PROYECTO.upper()); r.font.size = Pt(12); r.font.color.rgb = GRIS; r.bold = True
    p = doc.add_paragraph(); borde_parrafo(p, "bottom", "0F4C5C", "16")
    r = p.add_run(titulo); r.font.size = Pt(26); r.bold = True; r.font.color.rgb = PETROLEO
    p = doc.add_paragraph(); r = p.add_run(subtitulo); r.font.size = Pt(12); r.font.color.rgb = GRIS
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    p = doc.add_paragraph(); r = p.add_run("Contenido"); r.bold = True; r.font.size = Pt(14); r.font.color.rgb = PETROLEO
    campo(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

def convertir(md_path, salida, titulo_doc):
    lineas = md_path.read_text(encoding="utf-8").split("\n")
    doc = Document(); preparar(doc, titulo_doc); actualizar_campos_al_abrir(doc)
    # portada con el título y subtítulo del Markdown
    h1 = next((l[2:] for l in lineas if l.startswith("# ")), titulo_doc)
    h2 = next((l[3:] for l in lineas if l.startswith("## ")), "")
    portada(doc, h1, h2)
    saltar_titulo = {h1, h2}
    i, en_comentario = 0, False
    while i < len(lineas):
        l = lineas[i]
        s = l.strip()
        if "<!--" in s:
            en_comentario = "-->" not in s
            i += 1; continue
        if en_comentario:
            en_comentario = "-->" not in s; i += 1; continue
        if s.startswith("```"):
            lenguaje = s[3:].strip().lower()
            j = i + 1; bloque = []
            while j < len(lineas) and not lineas[j].strip().startswith("```"):
                bloque.append(lineas[j]); j += 1
            (mermaid if lenguaje == "mermaid" else lambda d, b: codigo(d, b, lenguaje))(doc, bloque)
            i = j + 1; continue
        if s.startswith("|") and i + 1 < len(lineas) and re.match(r"^\s*\|[\s:\-|]+\|\s*$", lineas[i + 1]):
            j = i; filas = []
            while j < len(lineas) and lineas[j].strip().startswith("|"):
                filas.append(lineas[j]); j += 1
            tabla(doc, filas); i = j; continue
        m = re.match(r"^(#{1,4})\s+(.*)$", l)
        if m:
            nivel, texto = len(m.group(1)), m.group(2).strip()
            if texto in saltar_titulo and nivel <= 2 and i < 5:
                i += 1; continue
            nivel = max(1, nivel - 1) if nivel >= 2 else 1   # ## → Título 1, ### → Título 2...
            h = doc.add_heading(level=nivel); agregar_texto(h, texto)
            i += 1; continue
        if s == "---":
            borde_parrafo(doc.add_paragraph(), "bottom", "CBD5E1", "6"); i += 1; continue
        if s.startswith(">"):
            p = doc.add_paragraph(); sombrear_parrafo(p, "F1F5F9"); borde_parrafo(p, "left", "0F4C5C", "18")
            p.paragraph_format.left_indent = Cm(0.3)
            agregar_texto(p, s.lstrip(">").strip(), Pt(9.5), GRIS); i += 1; continue
        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", l)
        if m:
            nivel = min(len(m.group(1)) // 2, 2)
            numerada = m.group(2)[0].isdigit()
            estilo = ("List Number" if numerada else "List Bullet") + (f" {nivel + 1}" if nivel else "")
            p = doc.add_paragraph(style=estilo); agregar_texto(p, m.group(3))
            i += 1; continue
        if s:
            agregar_texto(doc.add_paragraph(), s)
        i += 1
    doc.save(salida)

def resolver_mermaid(carpeta):
    """Devuelve (comando mmdc, puppeteer.json o None)."""
    if carpeta:
        base = pathlib.Path(carpeta).resolve()
        for nombre in ("mmdc.cmd", "mmdc"):
            exe = base / "node_modules" / ".bin" / nombre
            if exe.exists():
                pp = base / "puppeteer.json"
                return str(exe), (pp if pp.exists() else None)
        sys.exit(f"md_a_word.py: no encontré node_modules/.bin/mmdc en {base}")
    return shutil.which("mmdc"), None


def nombre_proyecto():
    srs = DOC / "SRS.md"
    if srs.exists():
        m = re.search(r"^## (?!\d|Anexo|Historial)(.+)$", srs.read_text(encoding="utf-8"), re.M)
        if m and "{{" not in m.group(1):
            return m.group(1).strip()
    return "Proyecto"


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Convierte los .md de documentacion/ a Word.")
    ap.add_argument("--mermaid", help="carpeta con mermaid-cli instalado (node_modules/.bin/mmdc)")
    ap.add_argument("--solo", help="convierte solo los documentos cuya ruta contiene este texto")
    ap.add_argument("--proyecto", help="nombre del proyecto (por defecto, del SRS)")
    a = ap.parse_args()
    MMDC, PUPPETEER = resolver_mermaid(a.mermaid)
    PROYECTO = a.proyecto or nombre_proyecto()
    SALIDA.mkdir(exist_ok=True)
    generados = 0
    for fuente, prefijo, titulo in documentos():
        if a.solo and a.solo not in fuente:
            continue
        destino = SALIDA / f"{prefijo}_{nombre_archivo(PROYECTO)}.docx"
        convertir(DOC / fuente, destino, titulo)
        generados += 1
        print("Generado:", destino.relative_to(DOC.parent))
    if not generados:
        print("No había documentos que convertir.")
